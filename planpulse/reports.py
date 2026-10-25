"""Weekly outputs: Excel KPI workbook, Word status note, PowerPoint market deck, static HTML dashboard."""
import csv
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C
from .narrative import read_approved

INK, TEAL, AMBER, PAPER = "1F2A33", "0F766E", "B45309", "F3EFE6"
DOCS = C.ROOT / "docs"


def gather(con, as_of: str) -> dict:
    tabs = {t[0] for t in con.execute("SHOW TABLES").fetchall()}
    g = dict(as_of=as_of)
    g["km"] = con.execute("SELECT * FROM kpi_market ORDER BY market").df()
    g["sku"] = con.execute("SELECT * FROM kpi_sku ORDER BY market, sku").df()
    g["acc"] = con.execute("SELECT * FROM kpi_accuracy_sku ORDER BY market, sku").df()
    g["otif"] = con.execute("SELECT * FROM kpi_otif_supplier ORDER BY market, supplier").df()
    g["exc"] = con.execute("SELECT * FROM exceptions ORDER BY severity, market, exc_type").df()
    g["act"] = con.execute("SELECT * FROM actions ORDER BY action_id").df()
    g["dq"] = con.execute("SELECT * FROM dq_results").df()
    g["bench"] = con.execute("SELECT * FROM ml_benchmark").df() if "ml_benchmark" in tabs else pd.DataFrame()
    g["late"] = con.execute("SELECT * FROM late_po ORDER BY days_late DESC").df()
    rw = con.execute("""SELECT k.market, k.sku, k.lead_time_days, k.stockout_risk, k.safety_stock,
                          list(p.end_balance ORDER BY p.wk) AS bal, list(p.demand ORDER BY p.wk) AS dem
                        FROM kpi_sku k JOIN proj p USING (market, sku)
                        WHERE k.stockout_risk <> 'LOW' GROUP BY ALL
                        ORDER BY k.stockout_risk, k.market, k.sku""").df()
    for c in ("bal", "dem"):
        rw[c] = rw[c].map(lambda v: [int(x) for x in v])
    g["runway"] = rw
    return g


def read_register(name: str) -> list[dict]:
    p = DOCS / name
    return list(csv.DictReader(open(p, encoding="utf-8"))) if p.exists() else []


# ---------------------------------------------------------------- Excel
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", fgColor=INK)
BODY = Font(name="Arial", size=10)


def _sheet(wb, title, df: pd.DataFrame, widths=None):
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) or v is pd.NaT else (v.date() if hasattr(v, "date") and not isinstance(v, (int, float, str)) else v) for v in r])
    _style(ws, widths)
    return ws


def _style(ws, widths=None):
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR_FONT, HDR_FILL, Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        w = (widths or {}).get(i) or min(46, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w


def build_xlsx(g: dict, path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    # --- Legend with the two parameters the formulas read
    lg = wb.create_sheet("Legend")
    rows = [
        ("PlanPulse KPI workbook", f"Week of {g['as_of']}. SIMULATED DATA, not real ISCP or company data."),
        ("EXCESS_DAYS", C.EXCESS_COVER_DAYS), ("Meaning", "Cover above this many days counts as excess (SKU_KPIs column R)"),
        ("", ""),
        ("Sheet", "What it holds"),
        ("Market_Summary", "Per-market KPIs. Every number is a formula over the other sheets."),
        ("SKU_KPIs", "One row per SKU with clean stock data. Columns I, J, N, O, R are formulas; the rest are values from the DuckDB KPI layer."),
        ("Accuracy_SKU", "Forecast accuracy for every SKU with sales and forecast, including SKUs held out of SKU_KPIs for stock data problems. Market WAPE and bias read from here."),
        ("OTIF_Supplier", "Closed POs in the trailing window. Percent columns are formulas."),
        ("Actions / Exceptions / DQ_Results", "Values exported from the tracker, exception list and data-quality checks."),
        ("", ""),
        ("Definitions", "Days of cover = on hand / (average weekly forecast, next 4 weeks / 7). WAPE = sum |forecast - actual| / sum actual. "
                        "Bias = (sum forecast - sum actual) / sum actual. OTIF = POs on time AND in full / POs received."),
        ("Stockout risk", "HIGH = projected stockout earlier than the lead time (a new order cannot land in time). MEDIUM = stockout within 8 weeks but after lead time. Value from SQL, not a formula."),
    ]
    for r in rows:
        lg.append(list(r))
    for row in lg.iter_rows():
        for c in row:
            c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical="top")
    lg["A1"].font = Font(name="Arial", bold=True, size=13, color=TEAL)
    lg["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    lg.column_dimensions["A"].width = 30
    lg.column_dimensions["B"].width = 110

    # --- SKU_KPIs
    s = g["sku"]
    ws = wb.create_sheet("SKU_KPIs")
    heads = ["market", "sku", "category", "on_hand", "safety_stock", "lead_time_days", "avg_weekly_fcst_4w", "unit_cost",
             "days_of_cover", "below_safety", "actual_units_8w", "forecast_units_8w", "abs_err_8w", "wape", "bias",
             "stockout_week", "stockout_risk", "excess_value"]
    ws.append(heads)

    def v(x):
        return None if pd.isna(x) else (x.item() if hasattr(x, "item") else x)
    for i, r in enumerate(s.itertuples(index=False), start=2):
        ws.append([r.market, r.sku, r.category, v(r.on_hand), v(r.safety_stock), v(r.lead_time_days), v(r.avg_weekly_fcst_4w), v(r.unit_cost),
                   f'=IF(AND(ISNUMBER(G{i}),G{i}>0),D{i}/(G{i}/7),"")',
                   f'=IF(AND(ISNUMBER(E{i}),D{i}<E{i}),1,0)',
                   v(r.actual_units), v(r.forecast_units_8w), v(r.abs_err),
                   f'=IF(AND(ISNUMBER(K{i}),K{i}>0),M{i}/K{i},"")',
                   f'=IF(AND(ISNUMBER(K{i}),K{i}>0),(L{i}-K{i})/K{i},"")',
                   v(r.stockout_week), r.stockout_risk,
                   f'=IF(ISNUMBER(G{i}),MAX(0,D{i}-Legend!$B$2*G{i}/7)*H{i},0)'])
    _style(ws)
    for col, fmt in {"I": "0.0", "N": "0.0%", "O": "0.0%", "R": "#,##0", "G": "0.0", "H": "#,##0.00"}.items():
        for c in ws[col][1:]:
            c.number_format = fmt
    fills = {"HIGH": "F4C7C3", "MEDIUM": "FCE8B2"}
    for c in ws["Q"][1:]:
        if c.value in fills:
            c.fill = PatternFill("solid", fgColor=fills[c.value])

    # --- Accuracy_SKU: every SKU with sales and forecast, including SKUs quarantined for stock data problems
    ac = g["acc"]
    wc = wb.create_sheet("Accuracy_SKU")
    wc.append(["market", "sku", "actual_units_8w", "forecast_units_8w", "abs_err_8w", "wape", "bias"])
    for i, r in enumerate(ac.itertuples(index=False), start=2):
        wc.append([r.market, r.sku, v(r.actual_units), v(r.forecast_units), v(r.abs_err),
                   f'=IF(C{i}>0,E{i}/C{i},"")', f'=IF(C{i}>0,(D{i}-C{i})/C{i},"")'])
    _style(wc)
    for col in "FG":
        for c in wc[col][1:]:
            c.number_format = "0.0%"

    # --- OTIF
    o = g["otif"]
    wo = wb.create_sheet("OTIF_Supplier")
    wo.append(["market", "supplier", "pos", "on_time", "in_full", "otif_pos", "on_time_pct", "in_full_pct", "otif_pct"])
    for i, r in enumerate(o.itertuples(index=False), start=2):
        wo.append([r.market, r.supplier, int(r.pos), int(r.on_time), int(r.in_full), int(r.otif_pos),
                   f"=D{i}/C{i}", f"=E{i}/C{i}", f"=F{i}/C{i}"])
    _style(wo)
    for col in "GHI":
        for c in wo[col][1:]:
            c.number_format = "0.0%"

    # --- value sheets
    act = g["act"][["action_id", "market", "severity", "exc_type", "ref", "owner_role", "description", "due_date", "status", "days_overdue",
                    "escalation_level", "created_week", "last_seen", "source", "note"]]
    _sheet(wb, "Actions", act, {7: 70, 15: 40})
    # Actions columns: B market, I status, J days_overdue (used by Market_Summary)
    _sheet(wb, "Exceptions", g["exc"][["market", "severity", "exc_type", "ref", "sku", "supplier", "detail", "owner_role"]], {7: 80})
    _sheet(wb, "DQ_Results", g["dq"][["market", "check_id", "severity", "description", "n_fail", "status"]], {4: 70})

    # --- Market_Summary (first sheet)
    ms = wb.create_sheet("Market_Summary", 0)
    ms["A1"] = f"PlanPulse market summary, week of {g['as_of']} (simulated data)"
    ms["A1"].font = Font(name="Arial", bold=True, size=13, color=TEAL)
    heads = ["Market", "SKUs", "Forecast WAPE", "Forecast bias", "Avg days of cover", "High stockout risk SKUs", "Below safety stock SKUs",
             "Excess value", "OTIF", "Open actions", "Overdue actions"]
    ms.append([])
    ms.append(heads)
    for c in ms[3]:
        c.font, c.fill, c.alignment = HDR_FONT, HDR_FILL, Alignment(wrap_text=True, vertical="center")
    mk = list(C.MARKETS)
    for k, m in enumerate(mk):
        r = 4 + k
        K = "SKU_KPIs!"
        ms.append([m,
                   f"=COUNTIFS({K}$A:$A,$A{r})",
                   f"=SUMIFS(Accuracy_SKU!$E:$E,Accuracy_SKU!$A:$A,$A{r})/SUMIFS(Accuracy_SKU!$C:$C,Accuracy_SKU!$A:$A,$A{r})",
                   f"=(SUMIFS(Accuracy_SKU!$D:$D,Accuracy_SKU!$A:$A,$A{r})-SUMIFS(Accuracy_SKU!$C:$C,Accuracy_SKU!$A:$A,$A{r}))/SUMIFS(Accuracy_SKU!$C:$C,Accuracy_SKU!$A:$A,$A{r})",
                   f"=AVERAGEIFS({K}$I:$I,{K}$A:$A,$A{r})",
                   f'=COUNTIFS({K}$A:$A,$A{r},{K}$Q:$Q,"HIGH")',
                   f"=SUMIFS({K}$J:$J,{K}$A:$A,$A{r})",
                   f"=SUMIFS({K}$R:$R,{K}$A:$A,$A{r})",
                   f"=SUMIFS(OTIF_Supplier!$F:$F,OTIF_Supplier!$A:$A,$A{r})/SUMIFS(OTIF_Supplier!$C:$C,OTIF_Supplier!$A:$A,$A{r})",
                   f'=COUNTIFS(Actions!$B:$B,$A{r},Actions!$I:$I,"Open")+COUNTIFS(Actions!$B:$B,$A{r},Actions!$I:$I,"In Progress")',
                   f'=COUNTIFS(Actions!$B:$B,$A{r},Actions!$J:$J,">0")'])
    t = 4 + len(mk)
    K = "SKU_KPIs!"
    ms.append(["All markets", f"=SUM(B4:B{t - 1})",
               "=SUM(Accuracy_SKU!$E:$E)/SUM(Accuracy_SKU!$C:$C)", "=(SUM(Accuracy_SKU!$D:$D)-SUM(Accuracy_SKU!$C:$C))/SUM(Accuracy_SKU!$C:$C)",
               f"=AVERAGE({K}$I:$I)", f"=SUM(F4:F{t - 1})", f"=SUM(G4:G{t - 1})", f"=SUM(H4:H{t - 1})",
               "=SUM(OTIF_Supplier!$F:$F)/SUM(OTIF_Supplier!$C:$C)", f"=SUM(J4:J{t - 1})", f"=SUM(K4:K{t - 1})"])
    for row in ms.iter_rows(min_row=4, max_row=t):
        for c in row:
            c.font = Font(name="Arial", size=10, bold=(c.row == t))
    for r in range(4, t + 1):
        ms[f"C{r}"].number_format = ms[f"D{r}"].number_format = ms[f"I{r}"].number_format = "0.0%"
        ms[f"E{r}"].number_format = "0.0"
        ms[f"H{r}"].number_format = "#,##0"
    ms.row_dimensions[3].height = 32
    for i, w in enumerate([16, 8, 14, 14, 16, 18, 18, 14, 10, 13, 14], 1):
        ms.column_dimensions[get_column_letter(i)].width = w
    ms.freeze_panes = "B4"
    ms["A10"] = "Excess value is in simulated currency units. Cover uses the next 4 weeks of system forecast."
    ms["A10"].font = Font(name="Arial", italic=True, size=9, color="666666")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- Word
def _top_exceptions(g, n=12):
    e = g["exc"]
    e = e[e.severity == "P1"].copy()
    return e.head(n)


def build_docx(g: dict, path: Path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    d = Document()
    for z in d.settings.element.xpath("w:zoom"):          # python-docx template omits the required percent attribute
        z.set(qn("w:percent"), "100")
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
        tcPr.append(sh)

    def heading(text, size=13):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(size), RGBColor.from_string(TEAL)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(3)

    def table(headers, rows, widths=None, fs=9):
        t = d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(h)
            run.bold, run.font.size, run.font.color.rgb = True, Pt(fs), RGBColor(255, 255, 255)
            shade(c, INK)
        for r in rows:
            cells = t.add_row().cells
            for i, val in enumerate(r):
                cells[i].text = ""
                cells[i].paragraphs[0].add_run(str(val)).font.size = Pt(fs)
        if widths:
            t.autofit = False
            for i, w in enumerate(widths):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        return t

    t = d.add_paragraph()
    r = t.add_run("PlanPulse weekly planning status"); r.bold, r.font.size, r.font.color.rgb = True, Pt(20), RGBColor.from_string(INK)
    p = d.add_paragraph(f"Week of {g['as_of']}. Markets: APAC, India, MENA-PAK, South Africa. All data is simulated and is not real ISCP or company data.")
    p.runs[0].italic = True

    km, act, exc = g["km"], g["act"], g["exc"]
    openA = act[act.status.isin(["Open", "In Progress"])]
    over = openA[openA.days_overdue > 0]
    heading("Headline")
    d.add_paragraph(f"{len(exc)} exceptions this week ({int((exc.severity == 'P1').sum())} at P1). {len(openA)} actions are open, {len(over)} are overdue and "
                    f"{int((over.escalation_level == 2).sum())} have reached the second escalation level. "
                    f"{int((g['dq'].n_fail > 0).sum())} market-level data-quality check failures.")

    heading("Market scorecard")
    table(["Market", "WAPE", "Bias", "Avg cover (d)", "High risk SKUs", "Late POs", "OTIF", "P1 exc."],
          [[r.market, f"{r.wape:.1%}", f"{r.bias:+.1%}", f"{r.avg_days_cover:.0f}", int(r.high_risk_skus), int(r.late_pos), f"{r.otif:.1%}",
            int((exc[(exc.market == r.market) & (exc.severity == 'P1')]).shape[0])] for r in km.itertuples()],
          [3.2, 1.8, 1.8, 2.4, 2.6, 1.8, 1.8, 1.6])

    heading("Exception summary")
    text, status = read_approved()
    if text:
        for para in text.split("\n\n"):
            d.add_paragraph(para.strip())
    d.add_paragraph(status).runs[0].italic = True

    heading("P1 exceptions needing action")
    rows = [[e.market, TYPE_NAME.get(e.exc_type, e.exc_type), e.ref, e.detail, e.owner_role] for e in _top_exceptions(g, 14).itertuples()]
    table(["Market", "Type", "Ref", "What is off", "Owner"], rows, [2.2, 2.6, 2.4, 6.2, 3.6], fs=8.5)
    n_p1 = int((exc.severity == "P1").sum())
    if n_p1 > 14:
        d.add_paragraph(f"{n_p1 - 14} more P1 exceptions are in the Excel workbook (Exceptions sheet).").runs[0].italic = True

    heading("Overdue and escalated actions")
    if len(over):
        top = over.sort_values(["escalation_level", "days_overdue"], ascending=False).head(10)
        table(["ID", "Market", "Sev", "Owner", "Due", "Days over", "Escalate to"],
              [[a.action_id, a.market, a.severity, a.owner_role, str(a.due_date)[:10], int(a.days_overdue), A_ESC.get(int(a.escalation_level), "")] for a in top.itertuples()],
              [1.8, 2.4, 1.1, 4.6, 2.2, 1.8, 2.6], fs=8.5)
    else:
        d.add_paragraph("No overdue actions this week.")

    heading("Data quality")
    fails = g["dq"][g["dq"].n_fail > 0]
    table(["Market", "Check", "Sev", "Rows"], [[f.market, f.check_id, f.severity, int(f.n_fail)] for f in fails.itertuples()], [3.4, 4, 2, 2], fs=8.5)

    heading("Top risks and process gaps")
    for rk in read_register("risk_register.csv")[:4]:
        d.add_paragraph(f"{rk['id']} ({rk['rating']}): {rk['risk']} Mitigation: {rk['mitigation']}", style="List Bullet")

    f = d.add_paragraph("Generated by PlanPulse. Figures come from the DuckDB KPI layer; definitions are in docs/glossary.md.")
    f.runs[0].font.size, f.runs[0].font.color.rgb = Pt(8.5), RGBColor(0x66, 0x66, 0x66)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


A_ESC = {1: "Market lead", 2: "ISCP lead"}
TYPE_NAME = {"LATE_PO": "Late PO", "STOCKOUT_RISK": "Stockout risk", "DATA_QUALITY": "Data quality", "FORECAST_BIAS": "Forecast bias",
             "BELOW_SAFETY": "Below safety", "EXCESS_STOCK": "Excess stock", "SUPPLIER_OTIF": "Supplier OTIF",
             "FORECAST_DIVERGENCE": "Forecast divergence", "UNEXPLAINED_ANOMALY": "Anomaly", "MEETING": "Meeting"}


# ---------------------------------------------------------------- PowerPoint
def build_pptx(g: dict, path: Path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    ink, teal, amber, paper = (RGBColor.from_string(x) for x in (INK, TEAL, AMBER, PAPER))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title, sub=None):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid(); s.background.fill.fore_color.rgb = paper
        tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
        tf = tb.text_frame; tf.word_wrap = True
        r = tf.paragraphs[0].add_run(); r.text = title
        r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(28), True, "Georgia", ink
        if sub:
            p = tf.add_paragraph(); rr = p.add_run(); rr.text = sub
            rr.font.size, rr.font.name, rr.font.color.rgb = Pt(13), "Calibri", RGBColor(0x55, 0x55, 0x55)
        foot = s.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12), Inches(0.3))
        fr = foot.text_frame.paragraphs[0].add_run()
        fr.text = f"PlanPulse | week of {g['as_of']} | simulated data, not real ISCP data"
        fr.font.size, fr.font.color.rgb, fr.font.name = Pt(10), RGBColor(0x77, 0x77, 0x77), "Calibri"
        return s

    def text(s, x, y, w, h, lines, size=14, bold_first=False, color=ink):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = ln
            r.font.size, r.font.name, r.font.color.rgb = Pt(size), "Calibri", color
            r.font.bold = bold_first and i == 0
            p.space_after = Pt(6)
        return tb

    def card(s, x, y, label, value, note=""):
        box = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.9), Inches(1.35))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box.line.color.rgb = RGBColor(0xD8, 0xD2, 0xC4)
        text(s, x + 0.12, y + 0.05, 2.7, 0.3, [label], 11, color=RGBColor(0x55, 0x55, 0x55))
        text(s, x + 0.12, y + 0.35, 2.7, 0.6, [value], 26, bold_first=True, color=teal)
        if note:
            text(s, x + 0.12, y + 0.92, 2.7, 0.4, [note], 10, color=RGBColor(0x66, 0x66, 0x66))

    def table(s, x, y, w, headers, rows, colw, fs=11, rowh=0.36):
        gt = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(rowh * (len(rows) + 1))).table
        for i, cw in enumerate(colw):
            gt.columns[i].width = Inches(cw)
        for i, h in enumerate(headers):
            c = gt.cell(0, i); c.text = h
            c.fill.solid(); c.fill.fore_color.rgb = ink
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.bold, r.font.color.rgb = Pt(fs), True, RGBColor(255, 255, 255)
        for ri, row in enumerate(rows, 1):
            for ci, val in enumerate(row):
                c = gt.cell(ri, ci); c.text = str(val)
                c.fill.solid(); c.fill.fore_color.rgb = RGBColor(255, 255, 255) if ri % 2 else RGBColor(0xF8, 0xF6, 0xF0)
                for p in c.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size, r.font.color.rgb = Pt(fs), ink
        return gt

    km, exc, act = g["km"], g["exc"], g["act"]
    openA = act[act.status.isin(["Open", "In Progress"])]
    over = openA[openA.days_overdue > 0]

    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = ink
    text(s, 0.8, 2.3, 11.5, 1.2, ["PlanPulse"], 54, bold_first=True, color=RGBColor(255, 255, 255))
    text(s, 0.8, 3.6, 11.5, 1.6, ["Weekly planning health: APAC, India, MENA-PAK, South Africa", f"Week of {g['as_of']}",
                                  "Simulated data. Not real ISCP or company data."], 18, color=RGBColor(0xE8, 0xE2, 0xD2))

    s = slide("Planning health by market", f"{len(exc)} exceptions, {int((exc.severity == 'P1').sum())} at P1; {len(over)} of {len(openA)} open actions are overdue")
    table(s, 0.6, 1.7, 6.9, ["Market", "WAPE", "Bias", "Cover (d)", "OTIF", "High risk"],
          [[r.market, f"{r.wape:.0%}", f"{r.bias:+.1%}", f"{r.avg_days_cover:.0f}", f"{r.otif:.0%}", int(r.high_risk_skus)] for r in km.itertuples()],
          [1.9, 0.9, 0.9, 1.0, 0.9, 1.3], 13, 0.5)
    cd = CategoryChartData(); cd.categories = list(km.market)
    cd.add_series("WAPE", [round(float(x), 3) for x in km.wape])
    cd.add_series("Absolute bias", [round(abs(float(x)), 3) for x in km.bias])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.8), Inches(1.6), Inches(5.0), Inches(4.6), cd)
    ch = gf.chart
    ch.has_title = True; ch.chart_title.text_frame.text = "Forecast error, last 8 weeks"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    ch.has_legend = True; ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.number_format, ch.plots[0].data_labels.number_format_is_linked = "0%", False
    ch.plots[0].data_labels.font.size = Pt(10)
    ch.series[0].format.fill.solid(); ch.series[0].format.fill.fore_color.rgb = teal
    ch.series[1].format.fill.solid(); ch.series[1].format.fill.fore_color.rgb = amber
    ch.value_axis.has_major_gridlines = False; ch.value_axis.tick_labels.font.size = Pt(10); ch.category_axis.tick_labels.font.size = Pt(10)
    ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = "0%", False
    text(s, 0.6, 4.6, 6.9, 1.8, ["WAPE is how far forecasts miss in absolute terms. Bias is the net direction of the miss.",
                                 "Cover is days of stock at the next 4 weeks of forecast. High risk means a stockout is projected sooner than the lead time."], 12,
         color=RGBColor(0x55, 0x55, 0x55))

    for r in km.itertuples():
        m = r.market
        s = slide(f"{m}: {int(r.high_risk_skus)} SKUs at high stockout risk, {int(r.late_pos)} late POs",
                  f"{int(r.skus)} SKUs in scope | {int(exc[exc.market == m].shape[0])} exceptions this week")
        card(s, 0.6, 1.7, "Forecast WAPE", f"{r.wape:.0%}", "last 8 weeks")
        card(s, 3.65, 1.7, "Forecast bias", f"{r.bias:+.1%}", "net direction")
        card(s, 0.6, 3.25, "Average cover", f"{r.avg_days_cover:.0f} d", "next 4 weeks of forecast")
        card(s, 3.65, 3.25, "OTIF", f"{r.otif:.0%}", "closed POs, trailing 26 weeks")
        p1 = exc[(exc.market == m) & (exc.severity == "P1")].head(5)
        lines = ["Needs action now (P1)"] + [f"{TYPE_NAME.get(e.exc_type, e.exc_type)}: {e.detail}" for e in p1.itertuples()] if len(p1) else ["Needs action now (P1)", "No P1 exceptions."]
        text(s, 7.0, 1.7, 5.8, 3.4, lines, 12, bold_first=True)
        dq = g["dq"][(g["dq"].market == m) & (g["dq"].n_fail > 0)]
        text(s, 0.6, 5.0, 12.2, 1.6, [f"Data quality: {len(dq)} failed check(s)" + (": " + ", ".join(dq.check_id) if len(dq) else ""),
                                      f"Overdue actions in this market: {int((over.market == m).sum())}"], 12, color=RGBColor(0x44, 0x44, 0x44))

    s = slide("Actions and escalations", f"{len(over)} overdue; escalation follows days past due against the severity SLA")
    top = over.sort_values(["escalation_level", "days_overdue"], ascending=False).head(9)
    table(s, 0.6, 1.7, 12.1, ["ID", "Market", "Sev", "Owner", "Due", "Days over", "Escalate to"],
          [[a.action_id, a.market, a.severity, a.owner_role, str(a.due_date)[:10], int(a.days_overdue), A_ESC.get(int(a.escalation_level), "")] for a in top.itertuples()]
          or [["", "", "", "No overdue actions", "", "", ""]],
          [1.2, 1.7, 0.8, 4.0, 1.6, 1.3, 1.5], 12, 0.45)

    s = slide("Risks and process gaps", "From docs/risk_register.csv and docs/gap_log.csv")
    risks = read_register("risk_register.csv")[:4]
    gaps = read_register("gap_log.csv")[:4]
    text(s, 0.6, 1.7, 6.0, 4.8, ["Risks"] + [f"{x['id']} ({x['rating']}): {x['risk']}" for x in risks], 12, bold_first=True)
    text(s, 6.9, 1.7, 5.9, 4.8, ["Process gaps"] + [f"{x['id']}: {x['gap']}" for x in gaps], 12, bold_first=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


# ---------------------------------------------------------------- HTML dashboard
HTML = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>PlanPulse dashboard</title><style>
:root{--ink:#1f2a33;--paper:#f3efe6;--teal:#0f766e;--amber:#b45309;--line:#d8d2c4;--card:#fff}
*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font:15px/1.45 "Avenir Next","Segoe UI Variable","Segoe UI",system-ui,sans-serif}
header{padding:22px 28px;background:var(--ink);color:#f3efe6}h1{margin:0;font:700 28px "Iowan Old Style","Palatino Linotype",Palatino,Georgia,serif}
header p{margin:4px 0 0;color:#cfc8b6;font-size:13px}main{padding:22px 28px;max-width:1300px;margin:auto}
nav{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:18px}nav button{border:1px solid var(--line);background:var(--card);padding:8px 14px;border-radius:4px;cursor:pointer;font:inherit}
nav button.on{background:var(--teal);color:#fff;border-color:var(--teal)}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin-bottom:20px}
.card{background:var(--card);border:1px solid var(--line);padding:12px 14px;border-radius:4px}.card small{color:#666;display:block}.card b{font:700 26px Georgia,serif;color:var(--teal)}
h2{font:700 18px Georgia,serif;margin:22px 0 8px}table{width:100%;border-collapse:collapse;background:var(--card);border:1px solid var(--line);font-size:13px}
th{background:var(--ink);color:#fff;text-align:left;padding:7px 9px;position:sticky;top:0;z-index:2}td{padding:6px 9px;border-top:1px solid #eee;vertical-align:top}
.P1{color:#9b1c1c;font-weight:700}.P2{color:var(--amber);font-weight:700}.P3{color:#555}.over{background:#fdf0e3}.wrap{overflow-x:auto;max-height:420px;overflow-y:auto}
select{font:inherit;padding:5px;margin-right:8px}.note{color:#666;font-size:12px;margin-top:20px}.mono{font-family:"Cascadia Mono","SF Mono",Consolas,monospace;font-variant-numeric:tabular-nums;font-size:12px}</style></head><body>
<header><h1>PlanPulse</h1><p>Week of __ASOF__ | simulated data, not real ISCP or company data | __SUMMARY__</p></header>
<main><nav id="nav"></nav><div id="cards" class="cards"></div>
<h2>Stockout runway, next 8 weeks</h2><p class="note" style="margin:0 0 8px">Each cell is projected end-of-week stock. Red is below zero, amber is under one week of demand. The outlined cell is the week a new order placed today would arrive: a red cell before it cannot be fixed by ordering, only by expediting.</p><div class="wrap"><table id="runway"></table></div>
<h2>Exceptions</h2><div><select id="sev"><option value="">All severities</option><option>P1</option><option>P2</option><option>P3</option></select>
<select id="typ"></select></div><div class="wrap"><table id="exc"></table></div>
<h2>Action tracker</h2><div class="wrap"><table id="act"></table></div>
<h2>Supplier OTIF</h2><div class="wrap"><table id="otif"></table></div>
<p class="note">Definitions: docs/glossary.md. Overdue rows are shaded. Escalation level 1 goes to the market lead, level 2 to the ISCP lead.</p></main>
<script>const D=__DATA__;let M="All";
const $=id=>document.getElementById(id);const pct=x=>x==null?"":(x*100).toFixed(1)+"%";
function rows(el,heads,data,cls){$(el).innerHTML="<tr>"+heads.map(h=>"<th>"+h[0]+"</th>").join("")+"</tr>"+data.map(r=>"<tr class='"+(cls?cls(r):"")+"'>"+heads.map(h=>"<td class='"+(h[2]?h[2](r):"")+"'>"+(h[1](r)??"")+"</td>").join("")+"</tr>").join("")}
function inM(r){return M==="All"||r.market===M}
function draw(){const km=D.km.filter(inM);const sum=k=>km.reduce((a,r)=>a+(r[k]||0),0);
$("cards").innerHTML=[["Forecast WAPE",pct(sum("abs_err_sum")/sum("actual_sum"))],["Forecast bias",((sum("forecast_sum")-sum("actual_sum"))/sum("actual_sum")*100).toFixed(1)+"%"],["Avg days of cover",(sum("cover_sum")/sum("cover_n")).toFixed(0)],
["High stockout risk SKUs",sum("high_risk_skus")],["Late POs",sum("late_pos")],["OTIF",pct(sum("otif_pos_sum")/sum("otif_n"))]].map(c=>"<div class='card'><small>"+c[0]+"</small><b>"+c[1]+"</b></div>").join("");
const sev=$("sev").value,typ=$("typ").value;
rows("exc",[["Sev",r=>r.severity,r=>r.severity],["Market",r=>r.market],["Type",r=>r.exc_type],["Ref",r=>r.ref],["What is off",r=>r.detail],["Owner",r=>r.owner_role]],
D.exc.filter(r=>inM(r)&&(!sev||r.severity===sev)&&(!typ||r.exc_type===typ)));
rows("act",[["ID",r=>r.action_id],["Sev",r=>r.severity,r=>r.severity],["Market",r=>r.market],["Owner",r=>r.owner_role],["Action",r=>r.description],["Due",r=>r.due_date],["Status",r=>r.status],
["Days over",r=>r.days_overdue],["Escalate to",r=>({1:"Market lead",2:"ISCP lead"})[r.escalation_level]||""]],D.act.filter(inM),r=>r.days_overdue>0?"over":"");
$("runway").innerHTML="<tr><th>SKU</th><th>Market</th><th>Risk</th><th>Lead time</th>"+[1,2,3,4,5,6,7,8].map(w=>"<th style='text-align:right'>W"+w+"</th>").join("")+"</tr>"+
D.runway.filter(inM).map(r=>{const lw=Math.floor(r.lead_time_days/7);return "<tr><td class='mono'>"+r.sku+"</td><td>"+r.market+"</td><td class='"+(r.stockout_risk==="HIGH"?"P1":"P2")+"'>"+r.stockout_risk+"</td><td>"+r.lead_time_days+" d</td>"+
r.bal.map((b,i)=>{const wk=r.dem[i]||1;const c=b<0?"#e8a39b":(b<wk?"#f3d38f":"#cfe3df");return "<td title='W"+(i+1)+": "+b+" units' class='mono' style='background:"+c+";text-align:right"+(i===lw?";outline:2px solid #1f2a33;outline-offset:-2px":"")+"'>"+b+"</td>"}).join("")+"</tr>"}).join("");
rows("otif",[["Market",r=>r.market],["Supplier",r=>r.supplier],["POs",r=>r.pos],["On time",r=>pct(r.on_time/r.pos)],["In full",r=>pct(r.in_full/r.pos)],["OTIF",r=>pct(r.otif_pos/r.pos)]],D.otif.filter(inM));}
$("nav").innerHTML=["All",...D.km.map(r=>r.market)].map(m=>"<button data-m='"+m+"' class='"+(m===M?"on":"")+"'>"+m+"</button>").join("");
$("nav").onclick=e=>{if(e.target.dataset.m){M=e.target.dataset.m;[...$("nav").children].forEach(b=>b.classList.toggle("on",b.dataset.m===M));draw()}};
$("typ").innerHTML="<option value=''>All types</option>"+[...new Set(D.exc.map(r=>r.exc_type))].sort().map(t=>"<option>"+t+"</option>").join("");
$("sev").onchange=$("typ").onchange=draw;draw();</script></body></html>"""


def build_html(g: dict, path: Path):
    def clean(df):
        df = df.copy()
        for c in df.columns:
            if str(df[c].dtype).startswith("datetime"):
                df[c] = df[c].dt.strftime("%Y-%m-%d")
        return json.loads(df.to_json(orient="records", date_format="iso"))
    data = dict(km=clean(g["km"]), exc=clean(g["exc"]), act=clean(g["act"]), otif=clean(g["otif"]), runway=clean(g["runway"]))
    act = g["act"]
    over = act[act.status.isin(["Open", "In Progress"]) & (act.days_overdue > 0)]
    html = (HTML.replace("__ASOF__", g["as_of"]).replace("__SUMMARY__", f"{len(g['exc'])} exceptions, {len(over)} overdue actions")
            .replace("__DATA__", json.dumps(data).replace("</", "<\\/")))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    return path


def write_all(g: dict, out_dir: Path):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    build_xlsx(g, out_dir / "PlanPulse_KPI_Workbook.xlsx")
    build_docx(g, out_dir / "PlanPulse_Weekly_Status_Note.docx")
    build_pptx(g, out_dir / "PlanPulse_Market_Status_Deck.pptx")
    build_html(g, out_dir / "dashboard.html")
    g["exc"].to_csv(out_dir / "exceptions.csv", index=False)
    g["act"].to_csv(out_dir / "action_tracker.csv", index=False)
